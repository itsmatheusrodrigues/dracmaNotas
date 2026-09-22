from pathlib import Path
from datetime import datetime
import re
import fitz  # PyMuPDF
from openpyxl import load_workbook


# ============================================================
# CONFIGURAÇÃO
# ============================================================
BASE_DIR = Path(__file__).resolve().parent
PASTA_NOTAS = BASE_DIR / "notas"
ARQUIVO_EXCEL = BASE_DIR / "Pasta1.xlsx"
ABA_EXCEL = "Planilha1"

# Preencha quando você definir os centros de custo de cada unidade consumidora.
# Exemplo:
# CENTRO_CUSTO_POR_UC = {
#     "3.392.699.013-61": "FAAMA",
#     "4.381.732.013-29": "RESERVATÓRIO DE ÁGUA",
# }
CENTRO_CUSTO_POR_UC = {}


# ============================================================
# UTILITÁRIOS
# ============================================================
def extrair_texto_pdf(caminho_pdf: Path) -> str:
    """Extrai o texto nativo de todas as páginas do PDF."""
    with fitz.open(caminho_pdf) as doc:
        return "\n".join(pagina.get_text("text") for pagina in doc)


def procurar(padrao: str, texto: str, grupo: int = 1, padrao_nao_encontrado: str = "") -> str:
    resultado = re.search(padrao, texto, flags=re.IGNORECASE | re.DOTALL)
    if not resultado:
        return padrao_nao_encontrado
    return resultado.group(grupo).strip()


def moeda_para_float(valor: str):
    if not valor:
        return None
    valor = valor.strip().replace("R$", "").replace(" ", "")
    valor = valor.replace(".", "").replace(",", ".")
    return float(valor)


def formatar_cnpj(cnpj: str) -> str:
    """Recebe um CNPJ com ou sem pontuação e devolve no formato 00.000.000/0000-00."""
    numeros = re.sub(r"\D", "", cnpj or "")
    if len(numeros) != 14:
        return ""

    return (
        f"{numeros[0:2]}.{numeros[2:5]}.{numeros[5:8]}/"
        f"{numeros[8:12]}-{numeros[12:14]}"
    )


def extrair_cnpj_emitente_por_chave(texto: str) -> str:
    """
    Extrai o CNPJ do emitente a partir da chave de acesso da NF-e/NF3e.

    Estrutura da chave:
    cUF(2) + AAMM(4) + CNPJ(14) + ...
    Portanto, o CNPJ ocupa as posições 7 a 20 da chave de 44 dígitos.
    """
    chave = procurar(
        r"chave de acesso:\s*([0-9\s\.\-]+)",
        texto,
    )

    chave_numerica = re.sub(r"\D", "", chave)

    if len(chave_numerica) >= 20:
        cnpj = chave_numerica[6:20]
        return formatar_cnpj(cnpj)

    return ""


def extrair_cnpj_generico(texto: str) -> str:
    """
    Tenta localizar um CNPJ completo no texto.
    É usado apenas como fallback quando a chave de acesso não fornece o CNPJ.
    """
    padrao_cnpj = r"\b(\d{2}\.\d{3}\.\d{3}/\d{4}-\d{2}|\d{14})\b"
    encontrados = re.findall(padrao_cnpj, texto)

    for cnpj in encontrados:
        formatado = formatar_cnpj(cnpj)
        if formatado:
            return formatado

    return ""


# ============================================================
# LEITOR - EQUATORIAL PARÁ
# ============================================================
def extrair_equatorial(caminho_pdf: Path, texto: str) -> dict:
    entidade = procurar(
        r"MODALIDADE TARIFÁRIA:.*?\n([A-ZÁÉÍÓÚÂÊÔÃÕÇ0-9 .&\-]+)\s*\nCNPJ:",
        texto,
    )

    # --------------------------------------------------------
    # EMITENTE = CNPJ DA EMPRESA
    # Primeiro tenta pela chave de acesso, que é mais confiável.
    # Se não encontrar, tenta localizar um CNPJ completo no texto.
    # --------------------------------------------------------
    emitente = extrair_cnpj_emitente_por_chave(texto)

    if not emitente:
        emitente = extrair_cnpj_generico(texto)

    # Como último fallback, mantém o nome do fornecedor para não deixar vazio.
    if not emitente:
        emitente = procurar(r"(EQUATORIAL PARÁ DISTRIB\. DE ENERGIA S\.A\.)", texto)

    numero_nf = procurar(r"NOTA FISCAL Nº\s*([0-9]+)", texto)

    data_emissao_txt = procurar(
        r"DATA DE EMISSÃO:\s*([0-9]{2}/[0-9]{2}/[0-9]{4})",
        texto,
    )
    data_emissao = (
        datetime.strptime(data_emissao_txt, "%d/%m/%Y")
        if data_emissao_txt
        else None
    )

    # Valor do documento no boleto.
    valor_txt = procurar(
        r"\(=\) VALOR DOCUMENTO\s*\n17\s*\nR\$\s*\n?\s*([0-9\.,]+)",
        texto,
    )
    if not valor_txt:
        valor_txt = procurar(
            r"R\$\s*([0-9]{1,3}(?:\.[0-9]{3})*,[0-9]{2})",
            texto,
        )
    valor = moeda_para_float(valor_txt)

    # Unidade consumidora + referência, logo após o beneficiário.
    unidade_consumidora = ""
    referencia = ""
    bloco_uc = re.search(
        r"EQUATORIAL PARÁ DISTRIB\. DE ENERGIA S\.A\.\s*\n"
        r"([0-9\.\-]+)\s*\n"
        r"([0-9]{2}/[0-9]{4})",
        texto,
        flags=re.IGNORECASE,
    )
    if bloco_uc:
        unidade_consumidora = bloco_uc.group(1).strip()
        referencia = bloco_uc.group(2).strip()

    linha_digitavel = procurar(
        r"([0-9]{5}\.[0-9]{5}\s+"
        r"[0-9]{5}\.[0-9]{6}\s+"
        r"[0-9]{5}\.[0-9]{6}\s+"
        r"[0-9]\s+[0-9]{14})",
        texto,
    )

    centro_custo = CENTRO_CUSTO_POR_UC.get(unidade_consumidora, "")

    return {
        "Entidade": entidade,
        "Nome do Documento": caminho_pdf.name,
        "Tipo de Documento": "Nota Fiscal de Energia Elétrica",
        "Emitente": emitente,
        "Data de Emissão": data_emissao,
        "Número": numero_nf,
        "Valor": valor,
        "Finalidade": (
            f"Pagamento de energia elétrica - referência {referencia}"
            if referencia
            else "Pagamento de energia elétrica"
        ),
        "Categorias de Gasto": "Energia Elétrica",
        "Centros de Custo": centro_custo,
        "Forma de Pagamento": "Boleto bancário / PIX",
        "Campo de Pagamento": linha_digitavel,
    }


# ============================================================
# CLASSIFICADOR DE DOCUMENTOS
# ============================================================
def analisar_documento(caminho_pdf: Path) -> dict:
    texto = extrair_texto_pdf(caminho_pdf)

    if "EQUATORIAL PARÁ DISTRIB" in texto.upper():
        return extrair_equatorial(caminho_pdf, texto)

    raise ValueError(
        "Modelo de documento ainda não cadastrado. "
        "Adicione um novo extrator para este fornecedor/tipo de documento."
    )


# ============================================================
# EXCEL
# ============================================================
def obter_mapa_colunas(ws) -> dict:
    mapa = {}
    for celula in ws[1]:
        if celula.value:
            mapa[str(celula.value).strip()] = celula.column
    return mapa


def gravar_documentos_no_excel(documentos: list[dict]):
    if not ARQUIVO_EXCEL.exists():
        raise FileNotFoundError(f"Excel não encontrado: {ARQUIVO_EXCEL}")

    wb = load_workbook(ARQUIVO_EXCEL)
    if ABA_EXCEL not in wb.sheetnames:
        raise ValueError(f"A aba '{ABA_EXCEL}' não existe no arquivo Excel.")

    ws = wb[ABA_EXCEL]
    colunas = obter_mapa_colunas(ws)

    colunas_obrigatorias = [
        "Entidade",
        "Nome do Documento",
        "Tipo de Documento",
        "Emitente",
        "Data de Emissão",
        "Número",
        "Valor",
        "Finalidade",
        "Categorias de Gasto",
        "Centros de Custo",
        "Forma de Pagamento",
        "Campo de Pagamento",
    ]

    faltando = [c for c in colunas_obrigatorias if c not in colunas]
    if faltando:
        raise ValueError(
            "O Excel não possui todas as colunas esperadas. Faltam: "
            + ", ".join(faltando)
        )

    # Evita duplicar um PDF que já foi processado.
    col_nome = colunas["Nome do Documento"]
    ja_processados = {
        str(ws.cell(row=linha, column=col_nome).value).strip()
        for linha in range(2, ws.max_row + 1)
        if ws.cell(row=linha, column=col_nome).value
    }

    adicionados = 0
    for doc in documentos:
        if doc["Nome do Documento"] in ja_processados:
            print(f"[IGNORADO] {doc['Nome do Documento']} já está no Excel.")
            continue

        linha = ws.max_row + 1
        for nome_coluna, valor in doc.items():
            ws.cell(row=linha, column=colunas[nome_coluna], value=valor)

        ws.cell(row=linha, column=colunas["Data de Emissão"]).number_format = "DD/MM/YYYY"
        ws.cell(row=linha, column=colunas["Valor"]).number_format = 'R$ #,##0.00'

        ja_processados.add(doc["Nome do Documento"])
        adicionados += 1
        print(f"[OK] {doc['Nome do Documento']} lançado na linha {linha}.")

    if adicionados:
        wb.save(ARQUIVO_EXCEL)
        print(f"\nExcel atualizado: {ARQUIVO_EXCEL}")
    else:
        print("\nNenhum documento novo para gravar.")


# ============================================================
# EXECUÇÃO
# ============================================================
def main():
    PASTA_NOTAS.mkdir(exist_ok=True)

    pdfs = sorted(PASTA_NOTAS.glob("*.pdf"))
    if not pdfs:
        print(f"Nenhum PDF encontrado em: {PASTA_NOTAS}")
        return

    documentos = []
    for pdf in pdfs:
        try:
            dados = analisar_documento(pdf)
            documentos.append(dados)

            if dados["Valor"] is not None:
                print(
                    f"[LIDO] {pdf.name}: NF {dados['Número']} | "
                    f"Emitente {dados['Emitente']} | "
                    f"R$ {dados['Valor']:.2f}"
                )
            else:
                print(
                    f"[LIDO] {pdf.name}: NF {dados['Número']} | "
                    f"Emitente {dados['Emitente']}"
                )

        except Exception as erro:
            print(f"[ERRO] {pdf.name}: {erro}")

    if documentos:
        gravar_documentos_no_excel(documentos)


if __name__ == "__main__":
    main()